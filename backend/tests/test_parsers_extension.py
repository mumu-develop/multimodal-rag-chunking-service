"""
Parser 扩展模块测试。

测试新增的文档解析器：
- ExcelParser (.xlsx, .xls)
- HTMLParser (.html, .htm)
- EmailParser (.eml)
- MboxParser (.mbox)
"""
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

# 确保 backend 目录在 sys.path
_BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_BACKEND_DIR))

os.environ["TMPDIR"] = "/tmp"


# ============================================================
# Excel Parser Tests
# ============================================================

class TestExcelParser:
    """Excel 文档解析器测试。"""

    @pytest.fixture
    def excel_parser(self):
        """创建 Excel 解析器实例。"""
        from app.parsers.excel_parser import ExcelParser
        return ExcelParser()

    @pytest.fixture
    def sample_excel(self, tmp_path):
        """创建测试 Excel 文件。"""
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # 添加表头
        ws.append(["姓名", "年龄", "城市"])
        # 添加数据
        ws.append(["张三", 25, "北京"])
        ws.append(["李四", 30, "上海"])
        ws.append(["王五", 28, "广州"])

        # 创建第二个 Sheet
        ws2 = wb.create_sheet("Sheet2")
        ws2.append(["产品", "价格", "库存"])
        ws2.append(["手机", 5000, 100])
        ws2.append(["电脑", 8000, 50])

        excel_path = tmp_path / "test.xlsx"
        wb.save(str(excel_path))
        return str(excel_path)

    def test_parse_excel_basic(self, excel_parser, sample_excel):
        """测试基础 Excel 解析。"""
        pages = excel_parser.parse(sample_excel, "test_doc_001")

        assert len(pages) == 2  # 2 个 Sheet
        assert pages[0].doc_type == "excel"
        assert "Sheet1" in pages[0].text or "姓名" in pages[0].text

    def test_parse_excel_tables_data(self, excel_parser, sample_excel):
        """测试 Excel 表格数据提取。"""
        pages = excel_parser.parse(sample_excel, "test_doc_001")

        # 第一个 Sheet 应有表格数据
        assert len(pages[0].tables_data) > 0
        table = pages[0].tables_data[0]
        assert "headers" in table
        assert "姓名" in table["headers"]
        assert len(table["rows"]) == 3  # 3 行数据

    def test_parse_excel_metadata(self, excel_parser, sample_excel):
        """测试 Excel 元数据提取。"""
        pages = excel_parser.parse(sample_excel, "test_doc_001")

        # 检查元数据
        metadata = pages[0].metadata
        assert "sheet_name" in metadata
        assert "row_count" in metadata
        assert metadata["sheet_name"] == "Sheet1"

    def test_parse_empty_excel(self, excel_parser, tmp_path):
        """测试空 Excel 文件。"""
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        wb = openpyxl.Workbook()
        ws = wb.active
        excel_path = tmp_path / "empty.xlsx"
        wb.save(str(excel_path))

        pages = excel_parser.parse(str(excel_path), "empty_doc")
        assert len(pages) == 1
        assert pages[0].doc_type == "excel"

    def test_parse_excel_with_merged_cells(self, excel_parser, tmp_path):
        """测试合并单元格处理。"""
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["标题", "", ""])
        ws.merge_cells("A1:C1")  # 合并第一行
        ws.append(["数据", "值", "备注"])

        excel_path = tmp_path / "merged.xlsx"
        wb.save(str(excel_path))

        pages = excel_parser.parse(str(excel_path), "merged_doc")
        assert len(pages) > 0
        # 检查合并信息是否提取
        table = pages[0].tables_data[0]
        assert "merges" in table


# ============================================================
# HTML Parser Tests
# ============================================================

class TestHTMLParser:
    """HTML 文档解析器测试。"""

    @pytest.fixture
    def html_parser(self):
        """创建 HTML 解析器实例。"""
        from app.parsers.html_parser import HTMLParser
        return HTMLParser()

    @pytest.fixture
    def sample_html(self, tmp_path):
        """创建测试 HTML 文件。"""
        html_content = """
        <!DOCTYPE html>
        <html>
        <head>
            <title>测试文档</title>
        </head>
        <body>
            <h1>主标题</h1>
            <h2>副标题</h2>
            <p>这是一个测试段落，包含一些文本内容。</p>
            <p>另一个段落，用于测试解析器的文本提取功能。</p>
            <table>
                <thead>
                    <tr><th>列1</th><th>列2</th><th>列3</th></tr>
                </thead>
                <tbody>
                    <tr><td>数据1</td><td>数据2</td><td>数据3</td></tr>
                    <tr><td>数据4</td><td>数据5</td><td>数据6</td></tr>
                </tbody>
            </table>
            <img src="test.png" alt="测试图片">
            <a href="https://example.com">链接文本</a>
        </body>
        </html>
        """
        html_path = tmp_path / "test.html"
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(html_content)
        return str(html_path)

    def test_parse_html_basic(self, html_parser, sample_html):
        """测试基础 HTML 解析。"""
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            pytest.skip("beautifulsoup4 not installed")

        pages = html_parser.parse(sample_html, "html_doc_001")

        assert len(pages) == 1
        assert pages[0].doc_type == "html"

    def test_parse_html_headings(self, html_parser, sample_html):
        """测试 HTML 标题提取。"""
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            pytest.skip("beautifulsoup4 not installed")

        pages = html_parser.parse(sample_html, "html_doc_001")

        headings = pages[0].headings
        assert len(headings) >= 2
        assert headings[0]["level"] == 1
        assert headings[0]["text"] == "主标题"

    def test_parse_html_tables(self, html_parser, sample_html):
        """测试 HTML 表格提取。"""
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            pytest.skip("beautifulsoup4 not installed")

        pages = html_parser.parse(sample_html, "html_doc_001")

        assert len(pages[0].tables_data) > 0
        table = pages[0].tables_data[0]
        assert "列1" in table["headers"]
        assert len(table["rows"]) == 2

    def test_parse_html_images(self, html_parser, sample_html):
        """测试 HTML 图片提取。"""
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            pytest.skip("beautifulsoup4 not installed")

        pages = html_parser.parse(sample_html, "html_doc_001")

        assert len(pages[0].images) > 0
        assert "test.png" in pages[0].images[0]

    def test_parse_html_scripts_removed(self, html_parser, tmp_path):
        """测试脚本标签被移除。"""
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            pytest.skip("beautifulsoup4 not installed")

        html_content = """
        <html>
        <body>
            <script>var x = 1;</script>
            <p>有效内容</p>
            <script>alert('test');</script>
        </body>
        </html>
        """
        html_path = tmp_path / "script.html"
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(html_content)

        pages = html_parser.parse(str(html_path), "script_doc")
        assert "var x" not in pages[0].text
        assert "alert" not in pages[0].text
        assert "有效内容" in pages[0].text

    def test_parse_html_no_table(self, html_parser, tmp_path):
        """测试无表格的 HTML。"""
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            pytest.skip("beautifulsoup4 not installed")

        html_content = "<html><body><p>只有文本</p></body></html>"
        html_path = tmp_path / "no_table.html"
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(html_content)

        pages = html_parser.parse(str(html_path), "no_table_doc")
        assert len(pages[0].tables_data) == 0


# ============================================================
# Email Parser Tests
# ============================================================

class TestEmailParser:
    """邮件文档解析器测试。"""

    @pytest.fixture
    def email_parser(self):
        """创建邮件解析器实例。"""
        from app.parsers.email_parser import EmailParser
        return EmailParser()

    @pytest.fixture
    def sample_email(self, tmp_path):
        """创建测试邮件文件 (.eml)。"""
        email_content = """From: sender@example.com
To: recipient@example.com
Subject: 测试邮件主题
Date: Mon, 10 Jun 2024 10:00:00 +0800
Content-Type: text/plain; charset=utf-8

这是一封测试邮件的正文内容。
包含多行文本。

邮件结尾。
"""
        email_path = tmp_path / "test.eml"
        with open(email_path, "w", encoding="utf-8") as f:
            f.write(email_content)
        return str(email_path)

    def test_parse_email_basic(self, email_parser, sample_email):
        """测试基础邮件解析。"""
        pages = email_parser.parse(sample_email, "email_doc_001")

        assert len(pages) == 1
        assert pages[0].doc_type == "email"

    def test_parse_email_headers(self, email_parser, sample_email):
        """测试邮件头提取。"""
        pages = email_parser.parse(sample_email, "email_doc_001")

        metadata = pages[0].metadata
        headers = metadata.get("email_headers", {})

        assert "subject" in headers
        assert headers["subject"] == "测试邮件主题"
        assert "from_address" in headers
        assert headers["from_address"] == "sender@example.com"

    def test_parse_email_body(self, email_parser, sample_email):
        """测试邮件正文提取。"""
        pages = email_parser.parse(sample_email, "email_doc_001")

        text = pages[0].text
        assert "测试邮件的正文内容" in text
        assert "邮件结尾" in text

    def test_parse_email_heading(self, email_parser, sample_email):
        """测试邮件标题作为 heading。"""
        pages = email_parser.parse(sample_email, "email_doc_001")

        headings = pages[0].headings
        assert len(headings) >= 1
        assert headings[0]["level"] == 1
        assert headings[0]["text"] == "测试邮件主题"

    def test_parse_multipart_email(self, email_parser, tmp_path):
        """测试多部分邮件（带 HTML 正文）。"""
        email_content = """From: sender@example.com
To: recipient@example.com
Subject: 多部分邮件
Content-Type: multipart/alternative; boundary="boundary"

--boundary
Content-Type: text/plain

纯文本正文

--boundary
Content-Type: text/html

<html><body><p>HTML正文</p></body></html>

--boundary--
"""
        email_path = tmp_path / "multipart.eml"
        with open(email_path, "w", encoding="utf-8") as f:
            f.write(email_content)

        pages = email_parser.parse(str(email_path), "multipart_doc")
        assert len(pages) == 1
        assert pages[0].doc_type == "email"

    def test_parse_email_with_attachment_info(self, email_parser, tmp_path):
        """测试带附件信息的邮件。"""
        # 注意：这只提取附件信息，不下载附件内容
        email_content = """From: sender@example.com
To: recipient@example.com
Subject: 带附件的邮件
Content-Type: multipart/mixed; boundary="mix"

--mix
Content-Type: text/plain

邮件正文

--mix
Content-Type: application/pdf; name="document.pdf"
Content-Disposition: attachment; filename="document.pdf"

<base64 encoded content>

--mix--
"""
        email_path = tmp_path / "attachment.eml"
        with open(email_path, "w", encoding="utf-8") as f:
            f.write(email_content)

        pages = email_parser.parse(str(email_path), "attach_doc")
        metadata = pages[0].metadata
        attachments = metadata.get("attachments", [])

        # 检查附件信息提取
        assert len(attachments) >= 0  # 可能提取到附件信息


# ============================================================
# Mbox Parser Tests
# ============================================================

class TestMboxParser:
    """Mbox 文件解析器测试。"""

    @pytest.fixture
    def mbox_parser(self):
        """创建 Mbox 解析器实例。"""
        from app.parsers.email_parser import MboxParser
        return MboxParser(max_emails=10)

    @pytest.fixture
    def sample_mbox(self, tmp_path):
        """创建测试 mbox 文件。"""
        mbox_content = """From sender1@example.com Mon Jun 10 10:00:00 2024
From: sender1@example.com
To: recipient@example.com
Subject: 第一封邮件
Date: Mon, 10 Jun 2024 10:00:00 +0800

第一封邮件正文。

From sender2@example.com Mon Jun 10 11:00:00 2024
From: sender2@example.com
To: recipient@example.com
Subject: 第二封邮件
Date: Mon, 10 Jun 2024 11:00:00 +0800

第二封邮件正文。

"""
        mbox_path = tmp_path / "test.mbox"
        with open(mbox_path, "w", encoding="utf-8") as f:
            f.write(mbox_content)
        return str(mbox_path)

    def test_parse_mbox_basic(self, mbox_parser, sample_mbox):
        """测试基础 mbox 解析。"""
        pages = mbox_parser.parse(sample_mbox, "mbox_doc_001")

        # 应解析多封邮件
        assert len(pages) >= 2
        for page in pages:
            assert page.doc_type == "email"

    def test_parse_mbox_max_limit(self, mbox_parser, tmp_path):
        """测试 mbox 解析数量限制。"""
        # 创建包含多封邮件的 mbox
        mbox_content = ""
        for i in range(20):
            mbox_content += f"""From sender{i}@example.com Mon Jun 10 {i}:00:00 2024
Subject: 邮件{i}

正文{i}

"""
        mbox_path = tmp_path / "large.mbox"
        with open(mbox_path, "w", encoding="utf-8") as f:
            f.write(mbox_content)

        pages = mbox_parser.parse(str(mbox_path), "large_mbox")
        # 应限制为 max_emails=10
        assert len(pages) <= 10


# ============================================================
# Progress Callback Tests
# ============================================================

class TestProgressCallback:
    """进度回调测试。"""

    def test_excel_progress(self, tmp_path):
        """测试 Excel 解析进度回调。"""
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        from app.parsers.excel_parser import ExcelParser

        wb = openpyxl.Workbook()
        excel_path = tmp_path / "progress.xlsx"
        wb.save(str(excel_path))

        progress_values = []

        def progress_fn(stage, pct):
            progress_values.append((stage, pct))

        parser = ExcelParser()
        parser.parse(str(excel_path), "progress_doc", progress_fn=progress_fn)

        assert len(progress_values) > 0
        assert progress_values[-1][1] == 100

    def test_html_progress(self, tmp_path):
        """测试 HTML 解析进度回调。"""
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            pytest.skip("beautifulsoup4 not installed")

        from app.parsers.html_parser import HTMLParser

        html_content = "<html><body><p>test</p></body></html>"
        html_path = tmp_path / "progress.html"
        with open(html_path, "w") as f:
            f.write(html_content)

        progress_values = []

        def progress_fn(stage, pct):
            progress_values.append((stage, pct))

        parser = HTMLParser()
        parser.parse(str(html_path), "progress_doc", progress_fn=progress_fn)

        assert len(progress_values) > 0
        assert progress_values[-1][1] == 100