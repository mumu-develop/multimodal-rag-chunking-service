"""PPTX 转 Markdown 转换器。

将 PowerPoint 幻灯片转换为 Markdown 格式：
- 幻灯片标题 → # 标题
- 正文内容 → 普通文本
- 列表项 → - 列表
"""
import os
import zipfile
from typing import Optional

import structlog

from app.config import settings
from app.converters import MarkdownConverter
from app.converters.table_utils import build_table_html

logger = structlog.get_logger()


class PptxConverter(MarkdownConverter):
    """PPTX 文件转 Markdown 转换器。"""

    def convert(self, file_path: str, doc_id: str) -> str:
        """将 PPTX 文件转换为 Markdown（兼容旧接口）。"""
        result = self.convert_full(file_path, doc_id)
        return result["markdown"]

    def convert_full(self, file_path: str, doc_id: str) -> dict:
        """
        将 PPTX 文件转换为结构化数据。

        参数：
            file_path: PPTX 文件路径
            doc_id: 文档 ID

        返回：
            {
                "markdown": str,           # Markdown 文本
                "tables": list[dict],      # 表格结构化数据（含 html 字段）
                "images": list[dict],      # 图片信息
                "code_blocks": list[dict], # 代码块信息
            }
        """
        from pptx import Presentation

        prs = Presentation(file_path)
        lines: list[str] = []
        tables_data: list[dict] = []
        code_blocks: list[dict] = []
        table_idx = 0

        for slide_idx, slide in enumerate(prs.slides, start=1):
            # 幻灯片分隔符
            lines.append(f"\n--- SLIDE {slide_idx} ---\n")

            # 处理幻灯片中的形状
            for shape in slide.shapes:
                # 处理表格
                if shape.has_table:
                    table_data = self._extract_pptx_table(shape.table, table_idx)
                    tables_data.append(table_data)
                    lines.append(f"\n[TABLE_{table_idx}]\n")
                    table_idx += 1
                    continue

                # 处理文本
                if not hasattr(shape, "text") or not shape.text.strip():
                    continue

                text = shape.text.strip()

                # 判断是否是标题（通常在幻灯片顶部）
                is_title = self._is_title_shape(shape, slide_idx, slide)

                if is_title:
                    lines.append(f"\n# {text}\n")
                else:
                    # 处理列表和正文
                    for line in text.split("\n"):
                        line = line.strip()
                        if not line:
                            continue
                        # 如果是列表项（以数字或符号开头）
                        if line[0] in ("•", "●", "○", "-") or line[0].isdigit():
                            lines.append(f"- {line.lstrip('•●○-').lstrip('0123456789.')}")
                        else:
                            lines.append(line)

        # 提取图片
        image_paths = self.extract_images(file_path, doc_id)
        images_data = [{"path": p, "index": i} for i, p in enumerate(image_paths)]

        return {
            "markdown": "\n".join(lines),
            "tables": tables_data,
            "images": images_data,
            "code_blocks": code_blocks,
        }

    def _extract_pptx_table(self, table, table_idx: int) -> dict:
        """
        提取 PPTX 表格结构化数据。

        参数：
            table: python-pptx Table 对象
            table_idx: 表格索引

        返回：
            {"index": int, "headers": list, "rows": list, "html": str, ...}
        """
        rows_data = []
        for row in table.rows:
            row_cells = []
            for cell in row.cells:
                cell_text = cell.text.strip().replace("\n", " ")
                row_cells.append(cell_text)
            rows_data.append(row_cells)

        if not rows_data:
            return {"index": table_idx, "headers": [], "rows": [], "merges": [], "html": "", "row_count": 0, "col_count": 0}

        # 第一行作为表头
        headers = rows_data[0]
        data_rows = rows_data[1:] if len(rows_data) > 1 else []

        # 构建 HTML（PPTX 表格通常无合并单元格）
        html_str = build_table_html(headers, data_rows)

        return {
            "index": table_idx,
            "headers": headers,
            "rows": data_rows,
            "merges": [],  # PPTX 表格合并单元格检测较复杂，暂不实现
            "html": html_str,
            "row_count": len(rows_data),
            "col_count": len(headers),
        }

    def extract_images(self, file_path: str, doc_id: str) -> list[str]:
        """从 PPTX 的 ZIP 结构中提取嵌入图片。"""
        output_dir = settings.IMAGE_OUTPUT_DIR
        os.makedirs(output_dir, exist_ok=True)

        image_paths: list[str] = []
        try:
            with zipfile.ZipFile(file_path) as zf:
                media_names = [
                    n for n in zf.namelist()
                    if n.startswith("ppt/media/") and not n.endswith("/")
                ]

                for idx, name in enumerate(sorted(media_names), start=1):
                    ext = os.path.splitext(name)[1] or ".png"
                    dest_name = f"{doc_id}_img_{idx - 1}{ext}"
                    dest_path = os.path.join(output_dir, dest_name)
                    with open(dest_path, "wb") as f:
                        f.write(zf.read(name))
                    image_paths.append(dest_path)
        except Exception:
            logger.exception("pptx_image_extraction_failed", doc_id=doc_id)

        return image_paths

    def _is_title_shape(self, shape, slide_idx: int, slide) -> bool:
        """判断形状是否是标题。

        简单策略：
        1. 检查形状类型（Title 类型）
        2. 如果是幻灯片的第一个文本框，当作标题
        """
        # 检查形状类型
        shape_type = shape.shape_type
        # Title 类型在 python-pptx 中有特定值
        if hasattr(shape, "text_frame"):
            # 检查是否是标题占位符
            if shape.placeholder_format and shape.placeholder_format.type == 1:  # Title
                return True

        # 备用：第一个文本框当作标题
        text_shapes = [s for s in slide.shapes if hasattr(s, "text") and s.text.strip()]
        if text_shapes and shape == text_shapes[0]:
            return True

        return False