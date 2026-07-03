"""Excel 文件转换器。

将 Excel 文件转换为 Markdown 格式，每个 Sheet 作为一个表格。

特点：
- 使用 openpyxl 解析
- 每个 Sheet 转换为独立的表格切片
- 支持合并单元格
- 表格统一生成 HTML 格式
"""
import os
import structlog

from app.config import settings
from app.converters import MarkdownConverter
from app.converters.models import ConverterOutput, ConversionError, TableData, ImageData, CodeBlockData
from app.converters.table_utils import build_table_html

logger = structlog.get_logger()


class XlsxConverter(MarkdownConverter):
    """Excel 文件转换器。"""

    def __init__(self, max_rows_per_sheet: int = 1000):
        self._max_rows_per_sheet = max_rows_per_sheet

    def convert(self, file_path: str, doc_id: str) -> str:
        """读取 Excel 文件内容（兼容旧接口）。"""
        result = self.convert_full(file_path, doc_id)
        return result.markdown

    def convert_full(self, file_path: str, doc_id: str) -> ConverterOutput:
        """读取 Excel 文件并提取结构化数据。

        Args:
            file_path: Excel 文件路径
            doc_id: 文档 ID

        Returns:
            ConverterOutput 模型

        Raises:
            ConversionError: 文件读取失败
        """
        try:
            import openpyxl
        except ImportError:
            logger.error("openpyxl_not_installed")
            raise ConversionError(
                "openpyxl 未安装，无法解析 Excel 文件",
                doc_type="xlsx",
                file_path=file_path
            )

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        except Exception as e:
            logger.exception("xlsx_file_read_failed", file_path=file_path, error=str(e))
            raise ConversionError(f"无法读取 Excel 文件: {e}", doc_type="xlsx", file_path=file_path)

        markdown_parts: list[str] = []
        tables_data: list[TableData] = []
        images_data: list[ImageData] = []
        code_blocks_data: list[CodeBlockData] = []

        total_sheets = len(wb.sheetnames)

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            sheet = wb[sheet_name]

            # 提取表格数据
            table_data = self._extract_sheet(sheet, sheet_name, sheet_idx)

            if table_data.is_valid():
                tables_data.append(table_data)

                # Markdown 中用标记引用表格
                markdown_parts.append(f"## {sheet_name}\n\n[TABLE_{sheet_idx}]\n\n")

                logger.info(
                    "xlsx_sheet_parsed",
                    sheet_name=sheet_name,
                    sheet_idx=sheet_idx,
                    rows=table_data.row_count,
                    cols=table_data.col_count,
                )

        wb.close()

        markdown_text = "\n".join(markdown_parts)

        logger.info(
            "xlsx_file_converted",
            file_path=file_path,
            doc_id=doc_id,
            sheets=total_sheets,
            tables=len(tables_data),
        )

        return ConverterOutput(
            markdown=markdown_text,
            tables=tables_data,
            images=images_data,
            code_blocks=code_blocks_data,
            doc_type="xlsx",
            page_count=total_sheets,
        )

    def _extract_sheet(self, sheet, sheet_name: str, sheet_idx: int) -> TableData:
        """提取单个 Sheet 数据。

        Args:
            sheet: openpyxl Worksheet 对象
            sheet_name: Sheet 名称
            sheet_idx: Sheet 索引

        Returns:
            TableData 模型
        """
        max_row = min(sheet.max_row, self._max_rows_per_sheet)
        max_col = sheet.max_column

        if max_row == 0 or max_col == 0:
            return TableData(
                index=sheet_idx,
                headers=[],
                rows=[],
                merges=[],
                html="",
                row_count=0,
                col_count=0,
            )

        # 提取合并单元格信息（转换为 table_utils 格式）
        merges = []
        for merge in sheet.merged_cells.ranges:
            merges.append({
                "row": merge.min_row - 1,  # table_utils 格式：row 从 0 开始
                "col": merge.min_col - 1,
                "rowspan": merge.max_row - merge.min_row + 1,
                "colspan": merge.max_col - merge.min_col + 1,
            })

        # 提取数据
        rows: list[list[str]] = []
        for row_idx in range(1, max_row + 1):
            row_data: list[str] = []
            for col_idx in range(1, max_col + 1):
                cell = sheet.cell(row=row_idx, column=col_idx)
                value = cell.value

                if value is None:
                    value = ""
                elif isinstance(value, bool):
                    value = "是" if value else "否"
                else:
                    value = str(value)

                row_data.append(value)

            rows.append(row_data)

        # 第一行作为表头
        headers = rows[0] if rows else []
        data_rows = rows[1:] if len(rows) > 1 else []

        # 使用统一的 HTML 生成器
        html_str = build_table_html(headers, data_rows, merges, class_name="chunk-table")

        return TableData(
            index=sheet_idx,
            headers=headers,
            rows=data_rows,
            merges=merges,
            html=html_str,
            row_count=len(rows),
            col_count=max_col,
        )

    def extract_images(self, file_path: str, doc_id: str) -> list[str]:
        """Excel 文件通常不包含嵌入式图片，返回空列表。"""
        return []